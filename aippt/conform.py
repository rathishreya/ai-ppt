"""Master-conformance rebuild: place preserved content into PLACEHOLDERS on real layouts.

This is the core of "follow the slide master": every output slide is built on a master
layout, and its text lives in the layout's placeholders (title / content / subtitle /
footer / date), so position, font family, and size are INHERITED from the master — not
hard-coded on free shapes. Run text is copied verbatim; the only character property we keep
is bold (emphasis), which does not break font/size inheritance.

Creativity (styling the layouts/master) is a separate, later step — and because slides now
follow the master, editing the master restyles the whole deck at once.
"""
from __future__ import annotations
import copy
from pptx import Presentation
from pptx.util import Emu, Pt

from .ir import DeckIR, ShapeIR, ParaIR, RunIR
from .brand import BrandSpec
from .classify import classify_slide, SlidePlan
from .designsystem import apply_theme
from pptx.enum.text import MSO_AUTO_SIZE

# placeholder idx constants (standard template)
CTR_TITLE = TITLE = 0
BODY = CONTENT = SUBTITLE = 1
CONTENT2 = 2
DATE = 10
FOOTER = 11
SLIDENUM = 12


def _ph_by_idx(slide):
    return {ph.placeholder_format.idx: ph for ph in slide.placeholders}


def _ensure_ph(slide, idx):
    """add_slide clones only title/body; clone footer/date/number from the layout on demand
    so their text is preserved AND stays placeholder-based (still follows the master)."""
    phs = _ph_by_idx(slide)
    if idx in phs:
        return phs[idx]
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.idx == idx:
            slide.shapes._spTree.append(copy.deepcopy(lph._element))
            return _ph_by_idx(slide).get(idx)
    return None


def _fill_paras(ph, paras: list[ParaIR]):
    """Fill a placeholder's text frame from IR paragraphs, preserving runs verbatim.
    No font/size/color set -> inherits from the layout/master (true conformance)."""
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # shrink to fit -> never overflow
    except Exception:
        pass
    first = True
    wrote = False
    for para in paras:
        runs = [r for r in para.runs if r.text != ""]
        if not runs:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        for r in runs:
            run = p.add_run()
            run.text = r.text
            if r.bold:
                run.font.bold = True
        wrote = True
    return wrote


def _fill_text(ph, text: str):
    tf = ph.text_frame
    tf.clear()
    tf.paragraphs[0].text = text


def _paras_of(shapes) -> list[ParaIR]:
    out = []
    for sh in shapes:
        if sh is None:
            continue
        for para in sh.paras:
            if para.text.strip():
                out.append(para)
    return out


def _headline_support(plan: SlidePlan):
    """Split all non-chrome text into a big headline + supporting text (both preserved)."""
    shapes = []
    for s in (plan.title, plan.eyebrow, plan.subtitle):
        if s is not None:
            shapes.append(s)
    shapes += plan.body
    seen, uniq = set(), []
    for s in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
        if id(s) not in seen:
            seen.add(id(s)); uniq.append(s)
    tmax = max((s.max_size_pt for s in uniq), default=0.0)
    head = [s for s in uniq if tmax > 0 and s.max_size_pt >= 0.55 * tmax]
    rest = [s for s in uniq if s not in head]
    head_para = ParaIR(runs=[RunIR(text=" ".join(s.text.replace("\n", " ") for s in head).strip())])
    support = _paras_of(rest)
    return head_para, support


def _chrome(slide, plan: SlidePlan):
    """Footer + wordmark into the footer/date placeholders so their text is preserved."""
    if plan.footers:
        ph = _ensure_ph(slide, FOOTER)
        if ph is not None:
            _fill_text(ph, "   ".join(s.text.replace("\n", " ") for s in plan.footers).strip())
    if plan.wordmark:
        ph = _ensure_ph(slide, DATE)
        if ph is not None:
            _fill_text(ph, plan.wordmark.text.strip())


def _add_table(slide, deck, shape: ShapeIR, top_in=3.6):
    grid = shape.table
    rows, cols = len(grid), max(len(r) for r in grid)
    from pptx.util import Inches
    gt = slide.shapes.add_table(rows, cols, Inches(0.62), Inches(top_in),
                                Inches(deck.width_in - 1.24), Inches(min(3.3, 0.4 * rows))).table
    for ri, row in enumerate(grid):
        for ci in range(cols):
            cell = gt.cell(ri, ci)
            runs = row[ci] if ci < len(row) else []
            tf = cell.text_frame
            p = tf.paragraphs[0]
            for r in runs:
                rr = p.add_run(); rr.text = r.text
                if ri == 0:
                    rr.font.bold = True


def conform(deck: DeckIR, brand: BrandSpec) -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(deck.width_emu)
    prs.slide_height = Emu(deck.height_emu)
    apply_theme(prs, brand)  # brand fonts/colors live in the master; slides inherit them
    L = {l.name: l for l in prs.slide_masters[0].slide_layouts}

    for s in deck.slides:
        plan = classify_slide(s, deck)

        if plan.stype == "title":
            slide = prs.slides.add_slide(L["Title Slide"])
            phs = _ph_by_idx(slide)
            head, support = _headline_support(plan)
            _fill_paras(phs[CTR_TITLE], [head])
            if SUBTITLE in phs:
                _fill_paras(phs[SUBTITLE], support)

        elif plan.stype in ("section", "closing"):
            slide = prs.slides.add_slide(L["Section Header"])
            phs = _ph_by_idx(slide)
            head, support = _headline_support(plan)
            _fill_paras(phs[TITLE], [head])
            if BODY in phs:
                _fill_paras(phs[BODY], support)

        else:  # content
            slide = prs.slides.add_slide(L["Title and Content"])
            phs = _ph_by_idx(slide)
            if plan.title and TITLE in phs:
                _fill_paras(phs[TITLE], plan.title.paras)
            body_shapes = ([plan.eyebrow] if plan.eyebrow else []) + \
                          ([plan.subtitle] if plan.subtitle else []) + plan.body
            body_paras = _paras_of(body_shapes)
            if CONTENT in phs and body_paras:
                _fill_paras(phs[CONTENT], body_paras)
            for tb in plan.tables:
                _add_table(slide, deck, tb)

        _chrome(slide, plan)

    return prs
