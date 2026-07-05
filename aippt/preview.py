"""Self-contained wireframe rasterizer for a .pptx (Pillow only; NOT LibreOffice).
Approximates layout/hierarchy/color so we can eyeball a blind build. Serif fonts are
substituted with Liberation Serif, sans with Liberation Sans — proportions ~match.
"""
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

IN = 914400
SCALE = 96 / 72  # render at 96 px per 72pt-inch -> 1280x720 for 13.33x7.5
PXIN = 96

SERIF = "/usr/share/fonts/truetype/liberation/LiberationSerif-{}.ttf"
SANS = "/usr/share/fonts/truetype/liberation/LiberationSans-{}.ttf"
_cache = {}


def font(name, size_px, bold, italic):
    serif = name and any(h in name.lower() for h in ("cambria", "times", "georgia", "serif", "garamond"))
    fam = SERIF if serif else SANS
    style = ("BoldItalic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular")
    key = (fam, style, size_px)
    if key not in _cache:
        try:
            _cache[key] = ImageFont.truetype(fam.format(style), size_px)
        except Exception:
            _cache[key] = ImageFont.truetype(SANS.format("Regular"), size_px)
    return _cache[key]


def px(emu):
    return int(emu / IN * PXIN)


def hexcol(s):
    s = str(s).lstrip("#")
    return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def shape_fill(sh):
    try:
        from pptx.enum.dml import MSO_FILL_TYPE
        if sh.fill.type == MSO_FILL_TYPE.SOLID:
            return hexcol(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def line_col(sh):
    try:
        c = sh.line.color
        if c is not None and c.type is not None:
            return hexcol(c.rgb)
    except Exception:
        pass
    return None


def run_color(r):
    try:
        c = r.font.color
        if c is not None and c.type is not None:
            return hexcol(c.rgb)
    except Exception:
        pass
    return (30, 34, 44)


def draw_text(draw, sh):
    tf = sh.text_frame
    x0, y0 = px(sh.left), px(sh.top)
    w = px(sh.width)
    anchor = str(getattr(tf, "vertical_anchor", None) or "")
    # measure total height first for middle anchor
    lines = []  # (segments:[(text,font,color)], height)
    for p in tf.paragraphs:
        segs = []
        maxsz = 12
        for r in p.runs:
            if not r.text:
                continue
            sz = int((r.font.size.pt if r.font.size else 13) * SCALE)
            maxsz = max(maxsz, sz)
            segs.append((r.text, font(r.font.name, sz, bool(r.font.bold), bool(r.font.italic)),
                         run_color(r)))
        if not segs:
            lines.append(([], int(8 * SCALE))); continue
        # wrap
        cur, curw = [], 0
        wrapped = []
        for text, fnt, col in segs:
            for word in _tok(text):
                ww = fnt.getlength(word)
                if curw + ww > w and cur:
                    wrapped.append(cur); cur, curw = [], 0
                cur.append((word, fnt, col)); curw += ww
        if cur:
            wrapped.append(cur)
        for wl in wrapped or [[]]:
            lines.append((wl, int(maxsz * 1.22)))
    total_h = sum(h for _, h in lines)
    y = y0
    if "MIDDLE" in anchor:
        y = y0 + max(0, (px(sh.height) - total_h) // 2)
    for segs, h in lines:
        x = x0
        for word, fnt, col in segs:
            draw.text((x, y), word, font=fnt, fill=col)
            x += fnt.getlength(word)
        y += h


def _tok(text):
    out, cur = [], ""
    for ch in text:
        cur += ch
        if ch == " ":
            out.append(cur); cur = ""
    if cur:
        out.append(cur)
    return out


def draw_table(draw, sh):
    t = sh.table
    x0, y0, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
    rows, cols = len(t.rows), len(t.columns)
    rh, cw = h // rows, w // cols
    for ri in range(rows):
        for ci in range(cols):
            cell = t.cell(ri, ci)
            cx, cy = x0 + ci * cw, y0 + ri * rh
            fc = shape_fill(cell)
            if fc:
                draw.rectangle([cx, cy, cx + cw, cy + rh], fill=fc)
            draw.rectangle([cx, cy, cx + cw, cy + rh], outline=(200, 205, 212))
            txt = cell.text_frame.text[:60]
            col = (255, 255, 255) if (fc and sum(fc) < 300) else (30, 34, 44)
            draw.text((cx + 5, cy + 4), txt, font=font("sans", int(11 * SCALE), ri == 0, False), fill=col)


def render(path, out_prefix, only=None):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    for i, slide in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        img = Image.new("RGB", (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    import io as _io
                    pic = Image.open(_io.BytesIO(sh.image.blob)).convert("RGBA")
                    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
                    pic = pic.resize((max(1, w), max(1, h)))
                    img.paste(pic, (x, y), pic)
                    continue
                if sh.shape_type == MSO_SHAPE_TYPE.TABLE or getattr(sh, "has_table", False):
                    draw_table(d, sh); continue
                fc = shape_fill(sh)
                if fc is not None:
                    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
                    rad = min(24, w // 12, h // 6)
                    d.rounded_rectangle([x, y, x + w, y + h], radius=rad, fill=fc)
                lc = line_col(sh)
                if lc is not None and sh.shape_type == MSO_SHAPE_TYPE.LINE:
                    d.line([px(sh.left), px(sh.top), px(sh.left) + px(sh.width),
                            px(sh.top) + px(sh.height)], fill=lc, width=2)
                elif lc is not None and fc is None:
                    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
                    d.ellipse([x, y, x + w, y + h], outline=lc, width=2)
                if sh.has_text_frame and sh.text_frame.text.strip():
                    draw_text(d, sh)
            except Exception as e:
                pass
        img.save(f"{out_prefix}-{i:02d}.png")
    print("rendered", out_prefix)


if __name__ == "__main__":
    only = [int(x) for x in sys.argv[3].split(",")] if len(sys.argv) > 3 else None
    render(sys.argv[1], sys.argv[2], only)
