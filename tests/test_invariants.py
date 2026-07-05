"""Golden invariant test: building the AI71 sample preserves slides & text.
Skips if the (gitignored) sample deck isn't present locally."""
import os
import pytest

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "samples", "AI71", "AI71_source.pptx")


@pytest.mark.skipif(not os.path.exists(SAMPLE), reason="sample deck not present (gitignored)")
def test_ai71_invariants(tmp_path):
    from aippt.pipeline import build
    from pptx import Presentation
    out = str(tmp_path / "out.pptx")
    res = build(SAMPLE, out, fail_closed=True)
    assert res.report.ok
    assert res.report.slide_count_src == res.report.slide_count_out == 26
    # every slide must follow the master: content lives in placeholders
    prs = Presentation(out)
    for i, slide in enumerate(prs.slides, 1):
        assert any(sh.is_placeholder for sh in slide.shapes), f"slide {i} has no placeholders"


def test_textguard_detects_drop():
    """textguard must FAIL when a word is removed."""
    from aippt.ir import DeckIR, SlideIR, ShapeIR, ParaIR, RunIR
    from aippt.textguard import verify
    a = DeckIR(slides=[SlideIR(index=0, shapes=[ShapeIR(kind="text",
        paras=[ParaIR(runs=[RunIR(text="hello world")])])])])
    b = DeckIR(slides=[SlideIR(index=0, shapes=[ShapeIR(kind="text",
        paras=[ParaIR(runs=[RunIR(text="hello")])])])])
    assert verify(a, a).ok
    assert not verify(a, b).ok
